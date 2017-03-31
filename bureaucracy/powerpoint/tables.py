from pptx.shapes.table import Table, _Cell

from bureaucracy.powerpoint.engines import BaseEngine

CONTEXT_KEY_FOR_TABLE = 'PPT_CURRENT_TABLE'


class TableContainer(object):
    def __init__(self, table: Table):
        self.table = table

    def render(self, engine: BaseEngine, context: dict):
        context[CONTEXT_KEY_FOR_TABLE] = self

        try:
            contents_of_first_cell = self[0, 0].text
        except IndexError:
            return

        engine.render(contents_of_first_cell, context)

    def __getitem__(self, items):
        return CellContainer(self.table.cell(*items))

    @property
    def row_count(self):
        return len(self.table.rows)

    @property
    def column_count(self):
        return len(self.table.columns)


class CellContainer(object):
    def __init__(self, cell: _Cell):
        self.cell = cell

    @property
    def text(self):
        return self.cell.text_frame.paragraphs[0].text

    @text.setter
    def text(self, text):
        self.cell.text = text